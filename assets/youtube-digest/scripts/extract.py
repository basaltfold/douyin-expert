#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
YouTube 提文案 — 输入 YouTube 链接/ID 提取视频文字稿（开源实现版）
=====================================================================

数据源：yt-dlp（https://github.com/yt-dlp/yt-dlp，Unlicense）
  - 从 YouTube 的 player API 取 automatic captions（json3 / vtt）
  - 不依赖任何付费转售 API，不需要 DATA_API_KEY / DATA_API_BASE

与 Qoder 原版的差异：
  - 原版 POST /story/api/youtube/transcript（需要用户自建付费服务）
  - 本版直接调 yt-dlp 的 Python API，纯本地、无 Key

Usage:
    python extract.py "https://www.youtube.com/watch?v=dQw4w9WgXcQ"
    python extract.py "dQw4w9WgXcQ" --language "zh,en"
    python extract.py "https://youtu.be/dQw4w9WgXcQ" --timestamp
    python extract.py "URL" --json
"""

import argparse
import json
import os
import re
import sys
import tempfile
import time
from datetime import datetime
from pathlib import Path

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

try:
    import yt_dlp
    HAS_YTDLP = True
except ImportError:
    HAS_YTDLP = False

try:
    import requests
    HAS_REQUESTS = True
except ImportError:
    HAS_REQUESTS = False

# ─── 配置 ────────────────────────────────────────────────────────────────────────
DEFAULT_OUTPUT_DIR = Path.home() / "Downloads" / "DouyinExpertYoutubeDigest"
DEFAULT_LANGUAGE = "zh,zh-Hans,zh-CN,zh-TW,en"
SUBTITLE_FORMAT = "json3/vtt/best"

# ─── 终端颜色 ─────────────────────────────────────────────────────────────────────
GREEN = "\033[92m"
YELLOW = "\033[93m"
RED = "\033[91m"
CYAN = "\033[96m"
BOLD = "\033[1m"
DIM = "\033[2m"
RESET = "\033[0m"


def info(msg):
    print(f"{GREEN}[OK]{RESET} {msg}")


def warn(msg):
    print(f"{YELLOW}[!]{RESET} {msg}")


def error(msg):
    print(f"{RED}[x]{RESET} {msg}")


def step(msg):
    print(f"{CYAN}[->]{RESET} {msg}")


# ─── 工具函数 ─────────────────────────────────────────────────────────────────────
def fmt_ts(seconds):
    """秒 -> MM:SS 或 HH:MM:SS"""
    try:
        seconds = int(float(seconds))
    except (ValueError, TypeError):
        return "00:00"
    h, rem = divmod(seconds, 3600)
    m, s = divmod(rem, 60)
    if h:
        return f"{h:02d}:{m:02d}:{s:02d}"
    return f"{m:02d}:{s:02d}"


def extract_video_id(raw):
    """从各种 YouTube URL 形态里抠出视频 ID；如果已经是 ID 就原样返回。"""
    raw = (raw or "").strip()
    if not raw:
        return ""
    if re.fullmatch(r"[A-Za-z0-9_-]{11}", raw):
        return raw
    patterns = [
        r"(?:v=|/v/|youtu\.be/|/embed/|/shorts/)([A-Za-z0-9_-]{11})",
        r"^([A-Za-z0-9_-]{11})$",
    ]
    for pattern in patterns:
        m = re.search(pattern, raw)
        if m:
            return m.group(1)
    return ""


def video_url_of(video_id):
    return f"https://www.youtube.com/watch?v={video_id}" if video_id else ""


# ─── 字幕解析 ─────────────────────────────────────────────────────────────────────
def parse_json3(path):
    """解析 YouTube json3 字幕格式 -> [{start, duration, text}]"""
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        data = json.load(f)
    segments = []
    for event in data.get("events", []):
        segs = event.get("segs") or []
        text = "".join((s.get("utf8") or "") for s in segs).strip()
        if not text or text == "\n":
            continue
        start_ms = event.get("tStartMs", 0)
        dur_ms = event.get("dDurationMs", 0)
        segments.append({
            "start": round(start_ms / 1000.0, 3),
            "duration": round(dur_ms / 1000.0, 3),
            "text": text.replace("\n", " ").strip(),
        })
    return segments


def parse_vtt(path):
    """解析 WebVTT 字幕 -> [{start, duration, text}]"""
    segments = []
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        lines = f.read().splitlines()

    ts_re = re.compile(r"(\d{1,2}):(\d{2}):(\d{2})[.,](\d{3})\s*-->\s*(\d{1,2}):(\d{2}):(\d{2})[.,](\d{3})")
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        m = ts_re.search(line)
        if m:
            start = int(m.group(1)) * 3600 + int(m.group(2)) * 60 + int(m.group(3)) + int(m.group(4)) / 1000.0
            end = int(m.group(5)) * 3600 + int(m.group(6)) * 60 + int(m.group(7)) + int(m.group(8)) / 1000.0
            i += 1
            buf = []
            while i < len(lines) and lines[i].strip():
                buf.append(re.sub(r"<[^>]+>", "", lines[i]).strip())
                i += 1
            text = " ".join(x for x in buf if x).strip()
            if text:
                segments.append({
                    "start": round(start, 3),
                    "duration": round(max(end - start, 0), 3),
                    "text": text,
                })
        else:
            i += 1
    return segments


def dedupe_segments(segments):
    """YouTube 自动字幕常有重复行，去掉相邻完全重复的文本。"""
    out = []
    last = None
    for seg in segments:
        if seg["text"] == last:
            continue
        out.append(seg)
        last = seg["text"]
    return out


# ─── yt-dlp 取字幕 ────────────────────────────────────────────────────────────────
def _probe_info(url):
    """只取元数据与可用字幕轨列表，不下载任何字幕。"""
    probe_opts = {
        "skip_download": True,
        "quiet": True,
        "no_warnings": True,
        "noprogress": True,
    }
    with yt_dlp.YoutubeDL(probe_opts) as ydl:
        return ydl.extract_info(url, download=False)


def _available_sub_langs(info):
    """返回可下载的字幕语言代码集合（手写字幕 + 自动字幕）。"""
    langs = set()
    for key in ("subtitles", "automatic_captions"):
        block = info.get(key) or {}
        if isinstance(block, dict):
            langs.update(block.keys())
    return langs


def _pick_lang(wanted, available):
    """按优先级从 wanted 里挑第一个可用语言。"""
    for lang in wanted:
        if lang in available:
            return lang
    # 模糊匹配：zh-Hans 命中 zh、zh-CN 等
    for lang in wanted:
        base = lang.split("-")[0]
        for cand in available:
            if cand == base or cand.startswith(base + "-"):
                return cand
    return None


def fetch_with_ytdlp(video_id, language, want_metadata=True):
    """
    用 yt-dlp 拉取字幕 + 元数据。

    策略：先探测可用字幕轨，再**一次只下一个语言**（避免多语言并发触发 429）。
    某个语言 429 就等 2 秒换下一个候选，最多试 3 次。

    返回 dict：
      {video_id, language, transcript:[...], metadata:{title,author,duration}}
    """
    url = video_url_of(video_id)
    wanted = [x.strip() for x in language.split(",") if x.strip()]

    # 1) 探测
    try:
        info_dict = _probe_info(url)
    except Exception as exc:  # noqa: BLE001
        warn(f"yt-dlp 探测失败: {exc}")
        return None

    meta = {}
    if want_metadata and info_dict:
        meta = {
            "title": info_dict.get("title") or "",
            "author": info_dict.get("channel") or info_dict.get("uploader") or "",
            "duration": info_dict.get("duration") or 0,
            "view_count": info_dict.get("view_count") or 0,
            "like_count": info_dict.get("like_count") or 0,
            "upload_date": info_dict.get("upload_date") or "",
        }

    available = _available_sub_langs(info_dict or {})
    if not available:
        warn("该视频没有任何字幕轨（含自动生成）")
        return None

    # 2) 挑语言，逐个尝试
    candidates = []
    first = _pick_lang(wanted, available)
    if first:
        candidates.append(first)
    for cand in sorted(available):
        if cand not in candidates and cand.split("-")[0] in {w.split("-")[0] for w in wanted}:
            candidates.append(cand)
    if not candidates:
        # 兜底：直接取 en，再不行取第一个可用的
        candidates = ["en"] if "en" in available else [sorted(available)[0]]

    last_error = None
    for lang in candidates[:3]:
        tmpdir = tempfile.mkdtemp(prefix="dyexpert_yt_")
        outtmpl = str(Path(tmpdir) / "sub")
        ydl_opts = {
            "skip_download": True,
            "writeautomaticsub": True,
            "writesubtitles": True,
            "subtitleslangs": [lang],
            "subtitlesformat": SUBTITLE_FORMAT,
            "outtmpl": outtmpl,
            "quiet": True,
            "no_warnings": True,
            "noprogress": True,
        }
        try:
            with yt_dlp.YoutubeDL(ydl_opts) as ydl:
                ydl.download([url])
        except Exception as exc:  # noqa: BLE001
            last_error = exc
            warn(f"语言 {lang} 下载失败: {maybe_429(exc)}")
            # 429 时退避，避免连续触发
            if "429" in str(exc):
                time.sleep(2.0)
            continue

        files = sorted(Path(tmpdir).glob("sub*.json3")) + sorted(Path(tmpdir).glob("sub*.vtt"))
        if not files:
            continue

        chosen = files[0]
        try:
            if chosen.suffix == ".json3":
                segments = parse_json3(chosen)
            else:
                segments = parse_vtt(chosen)
        except Exception as exc:  # noqa: BLE001
            warn(f"字幕解析失败: {exc}")
            continue

        if not segments:
            continue

        return {
            "video_id": video_id,
            "language": lang,
            "transcript": dedupe_segments(segments),
            "metadata": meta,
            "_tmpdir": tmpdir,
        }

    if last_error is not None:
        warn(f"所有候选语言都失败，最后一次错误: {maybe_429(last_error)}")
    return None


def maybe_429(exc):
    """把 429 错误翻译成可读提示。"""
    s = str(exc)
    if "429" in s or "Too Many Requests" in s:
        return "YouTube 限流 (429)，请隔几分钟再试"
    return s


# ─── 输出格式化 ────────────────────────────────────────────────────────────────────
def render_lines(data, with_timestamp=True):
    lines = []
    for seg in data.get("transcript") or []:
        text = (seg.get("text") or "").strip()
        if not text:
            continue
        if with_timestamp:
            lines.append(f"[{fmt_ts(seg.get('start', 0))}] {text}")
        else:
            lines.append(text)
    return lines


def total_duration_of(data):
    transcript = data.get("transcript") or []
    if transcript:
        last = transcript[-1]
        try:
            return float(last.get("start", 0)) + float(last.get("duration", 0))
        except (ValueError, TypeError):
            pass
    meta_dur = (data.get("metadata") or {}).get("duration")
    try:
        return float(meta_dur or 0)
    except (ValueError, TypeError):
        return 0


def build_markdown(data, lines, with_timestamp, raw_input):
    video_id = data.get("video_id", "")
    language = data.get("language", "")
    metadata = data.get("metadata") or {}
    title = metadata.get("title") or ""
    author = metadata.get("author") or ""

    md = ["# YouTube 视频文案", ""]
    md.append("| 项目 | 内容 |")
    md.append("|------|------|")
    if title:
        md.append(f"| 标题 | {title} |")
    if author:
        md.append(f"| 频道 | {author} |")
    md.append(f"| 视频ID | {video_id} |")
    md.append(f"| 视频链接 | {video_url_of(video_id) or raw_input} |")
    md.append(f"| 语言 | {language} |")
    md.append(f"| 片段数 | {len(data.get('transcript') or [])} |")
    md.append(f"| 总时长 | {fmt_ts(total_duration_of(data))} |")
    md.append(f"| 时间戳 | {'含' if with_timestamp else '不含'} |")
    md.append(f"| 数据来源 | yt-dlp 直连 YouTube（开源实现） |")
    md.append(f"| 提取时间 | {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} |")
    md.append("")
    md.append("---")
    md.append("")
    md.extend(lines)
    md.append("")
    return "\n".join(md)



# --- Excel 导出 -------------------------------------------------------------------
def export_excel(data, lines, fpath):
    """导出 xlsx：标题 / 时长 / 视频链接 / 文案内容"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except ImportError:
        return False

    video_id = data.get("video_id", "")
    metadata = data.get("metadata") or {}
    title = metadata.get("title") or video_id

    wb = Workbook()
    ws = wb.active
    ws.title = "视频文案"
    ws.append(["标题", "时长", "视频链接", "文案内容"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.append([title, fmt_ts(total_duration_of(data)), video_url_of(video_id), "\n".join(lines)])

    for col, width in {"A": 40, "B": 10, "C": 52, "D": 100}.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(fpath)
    return True


# ─── 主流程 ────────────────────────────────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(description="YouTube 提文案 — 输入 YouTube 链接/ID 提取视频文字稿（yt-dlp 开源实现）")
    parser.add_argument("video_url", help="YouTube 视频 URL（完整/短链）或直接给视频 ID")
    parser.add_argument("--language", default=DEFAULT_LANGUAGE,
                        help=f"字幕语言优先级，逗号分隔（默认 {DEFAULT_LANGUAGE}）")
    parser.add_argument("--timestamp", action="store_true", help="输出带 [MM:SS] 时间戳的文案")
    parser.add_argument("--no-metadata", action="store_true", help="不获取视频元数据")
    parser.add_argument("--json", action="store_true", help="终端输出原始 JSON")
    parser.add_argument("--excel", action="store_true", help="同时导出 Excel（.xlsx，需 openpyxl）")
    parser.add_argument("--no-translate", dest="translate", action="store_false",
                        help="保留字幕原文（本实现不做机器翻译，该参数仅为兼容保留）")
    parser.set_defaults(translate=True)
    parser.add_argument("--no-save", action="store_true", help="不保存 Markdown 文件")
    parser.add_argument("--output-dir", default=str(DEFAULT_OUTPUT_DIR), help=f"输出目录（默认 {DEFAULT_OUTPUT_DIR}）")
    args = parser.parse_args()

    if not HAS_YTDLP:
        error("缺少 yt-dlp，请先执行：python -m pip install -U yt-dlp")
        sys.exit(1)

    video_id = extract_video_id(args.video_url)
    if not video_id:
        error(f"无法识别 YouTube 视频 ID: {args.video_url}")
        sys.exit(1)

    step(f"正在提取字幕：{video_id}（语言优先级：{args.language}）")

    data = fetch_with_ytdlp(video_id, args.language, want_metadata=not args.no_metadata)
    if data:
        data.pop("_tmpdir", None)

    if data is None:
        error("提取失败：未拿到任何字幕轨")
        error("可能原因：视频无字幕 / 视频不存在或受限 / 需要登录 / 网络受限")
        sys.exit(1)

    if args.json:
        print(json.dumps(data, ensure_ascii=False, indent=2))

    metadata = data.get("metadata") or {}
    title = metadata.get("title") or ""
    author = metadata.get("author") or ""

    info(f"提取成功：视频 {video_id}")
    if title:
        print(f"  标题：{title}")
    if author:
        print(f"  频道：{author}")
    print(f"  语言：{data.get('language')}  | 片段数：{len(data.get('transcript') or [])}  | 总时长：{fmt_ts(total_duration_of(data))}")
    print(f"  链接：{video_url_of(video_id)}")

    lines = render_lines(data, args.timestamp)
    if not lines:
        warn("文字稿为空")
    else:
        print(f"\n{DIM}{'-' * 60}{RESET}")
        print(f"{BOLD}文案全文（{len(lines)} 段）：{RESET}\n")
        for line in lines:
            print(line)
        print(f"{DIM}{'-' * 60}{RESET}\n")

    if not args.no_save:
        out_dir = Path(args.output_dir).expanduser()
        out_dir.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        fpath = out_dir / f"{video_id or 'transcript'}_{stamp}.md"
        fpath.write_text(build_markdown(data, lines, args.timestamp, args.video_url), encoding="utf-8")
        info(f"文案已保存：{fpath}")

        if args.excel:
            xpath = out_dir / f"{video_id or 'transcript'}_{stamp}.xlsx"
            if export_excel(data, lines, xpath):
                info(f"Excel 已保存：{xpath}")
            else:
                warn("未安装 openpyxl，无法导出 Excel，请执行：python -m pip install openpyxl")


if __name__ == "__main__":
    main()
